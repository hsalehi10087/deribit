#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
پایش و پرستاری پایانه‌های فروش (POS) - تحلیل مشتریان

ورودی : هر تعداد فایل اکسل ماهانه (POS و/یا IPG)
خروجی : یک فایل اکسل گزارش (جدول‌ها + نمودار معدل و تراکنش)

اجرا:
    python pos_analysis.py file1.xlsx file2.xlsx ... -o pos_monitoring_report.xlsx

نکته: نوع فایل (POS / IPG) و دوره (مثلاً 1405/06) از عنوان هر فایل خوانده می‌شود.
"""
import argparse
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# ------------------------------------------------------------------
# تنظیمات و آستانه‌ها (هرکدام را بخواهید همین‌جا عوض کنید)
# ------------------------------------------------------------------
CONFIG = {
    "rial_per_toman": 10,                  # داده‌ها ریال فرض شده‌اند؛ هر تومان = ۱۰ ریال
    "min_avg_balance_toman": 30_000_000,   # حداقل معدل برای طرح وام صبا (۳۰ میلیون تومان)
    "drop_pct": 0.30,                      # افت کارکرد ≥ ۳۰٪ نسبت به دوره قبل
    "stable_band": 0.10,                   # تغییر بین ±۱۰٪ = ثابت
    "weak_terminal_ratio": 0.20,           # دستگاهی که کارکردش < ۲۰٪ دستگاه قوی‌تر است، ضعیف
    "top_n": 10,                           # تعداد مشتری در هر نمودار
    # وزن‌های نمره ریسک (جمع = ۱۰۰)
    "risk_weight_churn": 40,               # ریزش کامل
    "risk_weight_decline": 30,             # افت کارکرد
    "risk_weight_low_volume": 30,          # کارکرد در ۲۵٪ پایین مشتریان فعال
}

PERIOD_LABEL = "{}/{}"


# ------------------------------------------------------------------
# خواندن فایل‌ها
# ------------------------------------------------------------------
def clean_id(x):
    """کد مشتری/حساب را به رشته تمیز تبدیل می‌کند (حذف فاصله و اعشار)."""
    if pd.isna(x):
        return ""
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    out = str(x).strip().replace(" ", "")
    if out.isdigit() and len(out) == 10 and out.startswith("9"):
        out = "0" + out          # شماره موبایل بدون صفر اول در فایل‌ها ذخیره شده
    return out


def to_num(s):
    return pd.to_numeric(s, errors="coerce").fillna(0)


def find_col(df, name, last=False, prefix=False):
    """ستون را با نام دقیق (یا پیشوند) پیدا می‌کند. last=True یعنی آخرین تکرار."""
    idx = [i for i, c in enumerate(df.columns)
           if (str(c).startswith(name) if prefix else str(c).strip() == name)]
    if not idx:
        raise KeyError(f"ستون پیدا نشد: {name}")
    return df.iloc[:, idx[-1] if last else idx[0]]


def read_snapshot(path):
    """عنوان، نوع فایل (POS/IPG) و دوره را می‌خواند و جدول داده را برمی‌گرداند."""
    raw = pd.read_excel(path, header=None, dtype=object)
    title = str(raw.iloc[0].dropna().iloc[0]).strip()
    m = re.search(r"(\d{6})\s*$", title)
    if not m:
        raise ValueError(f"دوره در عنوان فایل پیدا نشد: {path}")
    period = PERIOD_LABEL.format(m.group(1)[:4], m.group(1)[4:])
    kind = "POS" if "POS" in title else ("IPG" if "IPG" in title else "?")

    is_header = raw.apply(lambda r: r.astype(str).str.strip().eq("نوع تراکنش").any(), axis=1)
    hdr = raw.index[is_header][0]
    df = raw.loc[hdr + 1:].reset_index(drop=True)
    df.columns = [str(c).strip() for c in raw.loc[hdr]]
    return kind, period, df


def parse_pos(df, period):
    out = pd.DataFrame({
        "period": period,
        "terminal": to_num(find_col(df, "کد ترمینال", last=True)).astype("int64"),
        "cust": find_col(df, "شماره تعریف مشتری").map(clean_id),
        "acct": find_col(df, "شماره حساب").map(clean_id),
        "store": find_col(df, "نام فروشگاه").fillna("").astype(str).str.strip(),
        "owner": (find_col(df, "نام").fillna("").astype(str) + " "
                  + find_col(df, "نام خانوادگی").fillna("").astype(str)).str.strip(),
        "mobile": find_col(df, "شماره تلفن همراه").map(clean_id),
        "guild": find_col(df, "نام صنف").fillna("").astype(str).str.strip(),
        "model": find_col(df, "مدل پایانه").fillna("").astype(str).str.strip(),
        "cnt": to_num(find_col(df, "تعداد تراکنشها")),
        "amt": to_num(find_col(df, "مبلغ تراکنش ها")),
        "bal": to_num(find_col(df, "میانگین یک ماهه حساب")),
    })
    return out[~out["cust"].isin(["", "0"])].reset_index(drop=True)  # "0" = کد نامعتبر/بی‌مشتری


def parse_ipg(df, period):
    out = pd.DataFrame({
        "period": period,
        "acct": find_col(df, "شماره حساب").map(clean_id),
        "owner": find_col(df, "نام دارنده حساب").fillna("").astype(str).str.strip(),
        "store": find_col(df, "نام فروشگاه").fillna("").astype(str).str.strip(),
        "cnt": to_num(find_col(df, "تعداد تراکنش موفق")),
        "amt": to_num(find_col(df, "مبلغ تراکنش")),
        "bal": to_num(find_col(df, "میانگین مانده حساب طی ماه", prefix=True)),
    })
    return out[out["acct"] != ""].reset_index(drop=True)


# ------------------------------------------------------------------
# تحلیل مشتریان POS
# ------------------------------------------------------------------
def trend_label(churn, g, band, drop):
    return np.select(
        [churn, g.isna(), g <= -drop, g <= -band, g >= band],
        ["ریزش (بدون تراکنش)", "جدید / بازگشته", "افت شدید", "رو به کاهش", "رو به رشد"],
        default="ثابت",
    )


def analyse_pos(pos, cfg):
    periods = sorted(pos["period"].unique())
    L, P = periods[-1], periods[-2]
    rial = cfg["rial_per_toman"]
    min_bal_rial = cfg["min_avg_balance_toman"] * rial
    drop, band = cfg["drop_pct"], cfg["stable_band"]

    # ---- تجمیع مشتری × دوره
    cp = pos.groupby(["cust", "period"]).agg(
        terminals=("terminal", "nunique"), cnt=("cnt", "sum"), amt=("amt", "sum")).reset_index()
    bal_cp = (pos.drop_duplicates(["cust", "period", "acct"])
              .groupby(["cust", "period"])["bal"].sum().reset_index())
    cp = cp.merge(bal_cp, on=["cust", "period"], how="left")

    def wide(col):
        return cp.pivot(index="cust", columns="period", values=col).reindex(columns=periods)

    present = wide("cnt").notna()
    terms = wide("terminals").fillna(0)
    cnt = wide("cnt").fillna(0)
    amt = wide("amt").fillna(0)
    bal = wide("bal").fillna(0)
    active = cnt > 0
    active_L = active[L]
    churn = (~active_L) & active.drop(columns=[L]).any(axis=1)

    # ---- تغییر کارکرد؛ فقط روی پایانه‌های مشترک دو دوره (چون فایل ۱۴۰۵/۰۶ یک پایانه به ازای هر مشتری دارد)
    tp = pos.groupby(["cust", "terminal", "period"]).agg(cnt=("cnt", "sum"), amt=("amt", "sum")).reset_index()
    a = tp[tp.period == P].set_index(["cust", "terminal"])[["cnt", "amt"]]
    b = tp[tp.period == L].set_index(["cust", "terminal"])[["cnt", "amt"]]
    common = a.index.intersection(b.index)
    ca = a.loc[common].groupby(level=0).sum()
    cb = b.loc[common].groupby(level=0).sum()

    g_amt = (amt[L] / amt[P].replace(0, np.nan) - 1)
    g_cnt = (cnt[L] / cnt[P].replace(0, np.nan) - 1)
    g_amt_c = (cb["amt"] / ca["amt"].replace(0, np.nan) - 1).reindex(amt.index)
    g_cnt_c = (cb["cnt"] / ca["cnt"].replace(0, np.nan) - 1).reindex(amt.index)
    g_amt = g_amt_c.combine_first(g_amt)
    g_cnt = g_cnt_c.combine_first(g_cnt)
    g_amt[churn] = -1.0
    g_cnt[churn] = -1.0

    trend = pd.Series(trend_label(churn, g_amt, band, drop), index=amt.index)

    # ---- معدل (موجودی) و روند آن
    def bal_info(i):
        vals = [bal.at[i, p] for p in periods if present.at[i, p] and bal.at[i, p] > 0]
        if not vals:
            return 0.0, False, 0.0
        up = len(vals) >= 2 and vals[-1] > vals[0] and vals[-1] >= vals[-2]
        return vals[-1], up, vals[0]

    info = [bal_info(i) for i in amt.index]
    bal_last = pd.Series([x[0] for x in info], index=amt.index)
    bal_up = pd.Series([x[1] for x in info], index=amt.index)
    bal_first = pd.Series([x[2] for x in info], index=amt.index)

    high_bal = bal_last >= min_bal_rial
    potential = high_bal & bal_up
    decline_flag = churn | (g_amt <= -drop).fillna(False)  # کاهش مبلغ کارکرد (نه صرفاً تعداد تراکنش)
    saba = potential & decline_flag

    # ---- ترمینال‌های ازدست‌رفته (فقط گزارش؛ در نمره ریسک لحاظ نمی‌شود)
    tP = tp[tp.period == P].set_index(["cust", "terminal"])
    tL = tp[tp.period == L].set_index(["cust", "terminal"])
    lost_idx = tP.index.difference(tL.index)
    lost_count = pd.Series(lost_idx.get_level_values(0)).value_counts().reindex(amt.index, fill_value=0)

    # ---- نمره ریسک (۰ تا ۱۰۰)
    active_amt_L = amt[L][active_L]
    q25 = active_amt_L.quantile(0.25)
    decl = np.clip(-g_amt.fillna(0), 0, 1)
    risk = (cfg["risk_weight_churn"] * churn.astype(int)
            + cfg["risk_weight_decline"] * decl
            + cfg["risk_weight_low_volume"] * (active_L & (amt[L] <= q25)).astype(int))
    risk = risk.clip(0, 100).round(0)

    # ---- برتر / ضعیف / پتانسیل (برای نمودارها و برچسب‌ها)
    top_idx = amt[L][active_L].sort_values(ascending=False).index[: cfg["top_n"]]
    weak_all = churn | trend.isin(["افت شدید"])
    weak_rank = pd.DataFrame({"risk": risk, "amtP": amt[P]})[weak_all]
    weak_top = weak_rank.sort_values(["risk", "amtP"], ascending=False).index[: cfg["top_n"]]
    pot_top = bal_last[potential].sort_values(ascending=False).index[: cfg["top_n"]]

    # ---- دسته‌بندی متنی
    tags = pd.DataFrame(False, index=amt.index, columns=[
        "برتر", "ضعیف", "کاندید صبا", "دارای پتانسیل", "ریزش", "چند پوز"])
    tags.loc[top_idx, "برتر"] = True
    tags["ضعیف"] = weak_all
    tags["کاندید صبا"] = saba
    tags["دارای پتانسیل"] = potential & ~saba
    tags["ریزش"] = churn
    tags["چند پوز"] = terms[L] >= 2
    labels = tags.apply(lambda r: "، ".join([c for c in tags.columns if r[c]]), axis=1)

    # ---- اطلاعات نمایشی (آخرین اطلاعات شناخته‌شده مشتری)
    names = pos.sort_values("period").groupby("cust")[["owner", "store", "guild", "mobile"]].last()
    names = names.reindex(amt.index).fillna("")

    last_active = pd.Series("", index=amt.index)
    for p in periods:
        last_active[active[p]] = p

    cust = pd.DataFrame({
        "کد مشتری": amt.index,
        "صاحب پایانه": names["owner"].values,
        "نام فروشگاه": names["store"].values,
        "صنف": names["guild"].values,
        "موبایل": names["mobile"].values,
        f"تعداد پایانه ({L})": terms[L].astype(int).values,
    }, index=amt.index)
    for p in periods:
        cust[f"تراکنش {p}"] = cnt[p].astype(int).values
        cust[f"مبلغ {p} (میلیون تومان)"] = (amt[p] / (rial * 1e6)).round(2).values
        cust[f"معدل {p} (میلیون تومان)"] = (bal[p] / (rial * 1e6)).round(2).values
    cust["تغییر مبلغ (پایانه‌های مشترک، %)"] = g_amt.round(4).values
    cust["تغییر تراکنش (پایانه‌های مشترک، %)"] = g_cnt.round(4).values
    cust["وضعیت روند"] = trend.values
    cust["آخرین دوره فعال"] = last_active.values
    cust["نمره ریسک (۰-۱۰۰)"] = risk.values
    cust["دسته‌ها"] = labels.values
    cust = cust.sort_values("نمره ریسک (۰-۱۰۰)", ascending=False).reset_index(drop=True)

    # ---- فهرست‌های اختصاصی
    M = lambda s: (s / (rial * 1e6)).round(2)
    churn_df = pd.DataFrame({
        "کد مشتری": amt.index[churn], "صاحب پایانه": names["owner"][churn].values,
        "نام فروشگاه": names["store"][churn].values, "موبایل": names["mobile"][churn].values,
        "آخرین دوره فعال": last_active[churn].values,
        f"تراکنش آخرین دوره ({P or ''})": cnt[P][churn].astype(int).values,
        f"مبلغ آخرین دوره (میلیون تومان)": M(amt[P][churn]).values,
        "معدل آخرین دوره (میلیون تومان)": M(bal_last[churn]).values,
        "نمره ریسک": risk[churn].values,
    }).sort_values("مبلغ آخرین دوره (میلیون تومان)", ascending=False)

    def bal_block(mask):
        df = pd.DataFrame({
            "کد مشتری": amt.index[mask], "صاحب پایانه": names["owner"][mask].values,
            "نام فروشگاه": names["store"][mask].values, "موبایل": names["mobile"][mask].values,
            f"معدل اول (میلیون تومان)": M(bal_first[mask]).values,
            f"معدل آخر (میلیون تومان)": M(bal_last[mask]).values,
            "تغییر مبلغ (%)": g_amt[mask].round(3).values,
            "تغییر تراکنش (%)": g_cnt[mask].round(3).values,
            f"مبلغ {L} (میلیون تومان)": M(amt[L][mask]).values,
            f"تراکنش {L}": cnt[L][mask].astype(int).values,
            "وضعیت روند": trend[mask].values,
            "تعداد پایانه": terms[L][mask].astype(int).values,
        })
        return df.sort_values("معدل آخر (میلیون تومان)", ascending=False)

    saba_df = bal_block(saba)
    potential_df = bal_block(potential)
    potential_df["کاندید صبا"] = np.where(saba[potential].values, "بله", "خیر")

    # ---- چند پوز: با آخرین دوره‌ای که جزئیات پایانه کامل دارد
    counts = pos.groupby("period").agg(c=("cust", "nunique"), t=("terminal", "nunique"))
    detail_p = [p for p in periods if counts.at[p, "t"] > counts.at[p, "c"]]
    D = detail_p[-1] if detail_p else L
    Dprev = [p for p in periods if p < D][-1] if any(p < D for p in periods) else None
    td = tp[tp.period == D].copy()
    td_multi = td.groupby("cust").terminal.transform("nunique") >= 2
    td = td[td_multi].copy()
    td_info = pos[pos.period == D].groupby(["cust", "terminal"]).agg(
        store=("store", "last"), model=("model", "last"), guild=("guild", "last")).reset_index()
    td = td.merge(td_info, on=["cust", "terminal"], how="left")
    if Dprev:
        prev = tp[tp.period == Dprev][["cust", "terminal", "amt"]].rename(columns={"amt": "amt_prev"})
        td = td.merge(prev, on=["cust", "terminal"], how="left")
    else:
        td["amt_prev"] = np.nan
    td["max_amt"] = td.groupby("cust").amt.transform("max")
    td["min_amt"] = td.groupby("cust").amt.transform("min")
    td["نسبت به قوی‌ترین"] = (td.amt / td.max_amt).round(3)
    td["ضعیف‌ترین"] = (td.amt == td.min_amt) & (td["نسبت به قوی‌ترین"] < cfg["weak_terminal_ratio"])
    td["تغییر نسبت به دوره قبل"] = np.where(td.amt_prev > 0, td.amt / td.amt_prev - 1, np.nan)
    td = td.sort_values(["cust", "amt"], ascending=[True, True])
    multi_df = pd.DataFrame({
        "کد مشتری": td.cust, "کد پایانه": td.terminal, "نام فروشگاه": td.store,
        "مدل پایانه": td.model, "تراکنش": td.cnt.astype(int),
        f"مبلغ (میلیون تومان)": M(td.amt), f"مبلغ {Dprev or '-'} (میلیون تومان)": M(td.amt_prev),
        "تغییر نسبت به دوره قبل (%)": td["تغییر نسبت به دوره قبل"].round(3),
        "نسبت به قوی‌ترین پایانه": td["نسبت به قوی‌ترین"],
        "وضعیت": np.where(td["ضعیف‌ترین"], "ضعیف‌ترین دستگاه مشتری", ""),
    }).reset_index(drop=True)

    # ---- وضعیت ماهانه هر پایانه (غایب / فعال در هر دوره)
    tw = tp.pivot_table(index=["cust", "terminal"], columns="period", values="amt",
                        aggfunc="sum").reindex(columns=periods)
    tw_info = (pos.sort_values("period").groupby(["cust", "terminal"])[["store", "model"]]
               .last().reindex(tw.index))
    seen = tw.notna()
    term_last_seen = [next((p for p in reversed(periods) if seen.at[k, p]), "") for k in tw.index]
    term_missing_L = (seen.any(axis=1) & ~seen[L]).values
    term_df = pd.DataFrame({
        "کد مشتری": [k[0] for k in tw.index],
        "کد پایانه": [k[1] for k in tw.index],
        "نام فروشگاه": tw_info["store"].fillna("").values,
        "مدل پایانه": tw_info["model"].fillna("").values,
        "آخرین دوره دیده‌شده": term_last_seen,
        f"وضعیت در {L}": np.where(seen[L].values, "فعال", "غایب"),
    })
    for p in periods:
        term_df[f"مبلغ {p} (میلیون تومان)"] = (tw[p].fillna(0) / (rial * 1e6)).round(2).values
    term_df = term_df.sort_values(f"وضعیت در {L}").reset_index(drop=True)

    # ---- فهرست تماس اولویت‌دار
    miss_map = {}
    for k, is_missing in zip(tw.index, term_missing_L):
        if is_missing:
            miss_map.setdefault(k[0], []).append(str(k[1]))
    miss_count = pd.Series({c: len(v) for c, v in miss_map.items()}).reindex(amt.index, fill_value=0)

    months_out = {}
    for i in amt.index:
        last = next((p for p in reversed(periods) if active.at[i, p]), None)
        months_out[i] = periods.index(L) - periods.index(last) if last else np.nan
    months_out = pd.Series(months_out)
    last_active_s = pd.Series({i: next((p for p in reversed(periods) if active.at[i, p]), "")
                               for i in amt.index})

    value_m = amt.max(axis=1) / (rial * 1e6)
    has_bal = bal_last > 0
    reasons = pd.DataFrame({
        "خرابی دستگاه احتمالی": active_L & (miss_count > 0),
        "قهر/ترک احتمالی": churn & has_bal,
        "افت کارکرد": active_L & (g_amt <= -drop).fillna(False),
        "خروج کامل احتمالی": churn & ~has_bal,
    }, index=amt.index)
    weights = {"خرابی دستگاه احتمالی": 1.0, "قهر/ترک احتمالی": 0.9,
               "افت کارکرد": 0.7, "خروج کامل احتمالی": 0.4}
    action_map = {
        "خرابی دستگاه احتمالی": "بررسی فنی پایانه‌های غایب (احتمال خرابی) و تماس برای رفع مشکل",
        "قهر/ترک احتمالی": "تماس برای بازگشت مشتری و رفع نارضایتی (حساب هنوز موجودی دارد)",
        "افت کارکرد": "تماس برای بررسی علت کاهش کارکرد",
        "خروج کامل احتمالی": "تماس برای تأیید ترک و بررسی وضعیت حساب",
    }
    rw = pd.Series([max([weights[c] for c in reasons.columns if reasons.at[i, c]], default=0)
                    for i in amt.index], index=amt.index)
    cand = rw > 0
    vpct = value_m[cand].rank(pct=True)
    call_rows = []
    for i in amt.index[cand.values]:
        rs = [c for c in reasons.columns if reasons.at[i, c]]
        primary = max(rs, key=lambda c: weights[c])
        call_rows.append({
            "امتیاز اولویت (۰-۱۰۰)": round(60 * vpct[i] + 40 * rw[i]),
            "اقدام پیشنهادی": action_map[primary],
            "علت‌های احتمالی": "، ".join(rs),
            "کد مشتری": i,
            "صاحب پایانه": names.at[i, "owner"],
            "نام فروشگاه": names.at[i, "store"],
            "موبایل": names.at[i, "mobile"],
            "آخرین دوره فعال": last_active_s[i],
            "دوره‌های غیرفعال (تعداد)": int(months_out[i]) if not pd.isna(months_out[i]) else "",
            "پایانه‌های غایب": "، ".join(miss_map.get(i, [])),
            "حداکثر مبلغ ماهانه (میلیون تومان)": round(value_m[i], 2),
            "معدل آخر (میلیون تومان)": M(bal_last[i]),
            f"مبلغ {L} (میلیون تومان)": M(amt[L][i]),
            "وضعیت روند": trend[i],
        })
    call_df = pd.DataFrame(call_rows).sort_values("امتیاز اولویت (۰-۱۰۰)", ascending=False).reset_index(drop=True)
    call_df.insert(1, "موج تماس", np.where(call_df["امتیاز اولویت (۰-۱۰۰)"] >= 70, "موج ۱ (فوری)",
                   np.where(call_df["امتیاز اولویت (۰-۱۰۰)"] >= 40, "موج ۲", "موج ۳")))

    # ---- نمودارها: برتر، ضعیف، پتانسیل
    def chart_pack(idx):
        return {
            "custs": [f"{names.at[i, 'store'] or names.at[i, 'owner']} ({i})" for i in idx],
            "bal": [[round(bal.at[i, p] / (rial * 1e6), 2) for p in periods] for i in idx],
            "cnt": [[int(cnt.at[i, p]) for p in periods] for i in idx],
        }

    charts = {
        "مشتریان برتر": chart_pack(top_idx),
        "مشتریان ضعیف": chart_pack(weak_top),
        "مشتریان دارای پتانسیل": chart_pack(pot_top),
    }

    summary = {
        "periods": periods,
        "L": L, "P": P, "D": D, "Dprev": Dprev,
        "rows": [
            ("تعداد کل مشتریان (همه دوره‌ها)", len(amt)),
            (f"مشتریان دارای تراکنش در {L}", int(active_L.sum())),
            (f"ریزش (تراکنش قبلی، بدون تراکنش در {L})", int(churn.sum())),
            ("مشتریان ضعیف (ریزش + افت شدید ≥۳۰٪)", int(weak_all.sum())),
            ("کاندید وام صبا", int(saba.sum())),
            ("مشتریان دارای پتانسیل (معدل بالا و افزایشی)", int(potential.sum())),
            (f"مشتریان با بیش از یک پایانه ({D})", int(len(td.cust.unique()))),
            (f"دستگاه‌های ضعیف‌تر در مشتریان چندپوز ({D})", int(td["ضعیف‌ترین"].sum())),
            (f"پایانه‌های قطع‌شده بین {P} و {L} (گزارش، بدون لحاظ در ریسک)", int(len(lost_idx))),
        ],
        "period_stats": counts,
    }
    summary["rows"] += [
        ("فهرست تماس اولویت‌دار (تعداد کل)", len(call_df)),
        ("  ─ خرابی دستگاه احتمالی", int(reasons["خرابی دستگاه احتمالی"].sum())),
        ("  ─ قهر/ترک احتمالی (حساب موجودی دارد)", int(reasons["قهر/ترک احتمالی"].sum())),
        ("  ─ افت کارکرد ≥۳۰٪", int(reasons["افت کارکرد"].sum())),
        ("  ─ خروج کامل احتمالی", int(reasons["خروج کامل احتمالی"].sum())),
    ]
    return {
        "summary": summary, "customers": cust, "churn": churn_df, "saba": saba_df,
        "potential": potential_df, "multi": multi_df, "charts": charts,
        "calls": call_df, "terminals": term_df,
    }


# ------------------------------------------------------------------
# تحلیل درگاه IPG (سطح حساب)
# ------------------------------------------------------------------
def analyse_ipg(ipg, cfg):
    periods = sorted(ipg["period"].unique())
    rial = cfg["rial_per_toman"]
    g = ipg.groupby(["acct", "period"]).agg(cnt=("cnt", "sum"), amt=("amt", "sum"),
                                            bal=("bal", "max")).reset_index()
    names = ipg.sort_values("period").groupby("acct")[["owner", "store"]].last()
    W = lambda c: g.pivot(index="acct", columns="period", values=c).reindex(columns=periods).fillna(0)
    cnt, amt, bal = W("cnt"), W("amt"), W("bal")
    out = pd.DataFrame({"کد حساب": amt.index, "دارنده حساب": names["owner"].reindex(amt.index).values,
                        "نام فروشگاه": names["store"].reindex(amt.index).values}, index=amt.index)
    for p in periods:
        out[f"تراکنش {p}"] = cnt[p].astype(int).values
        out[f"مبلغ {p} (میلیون تومان)"] = (amt[p] / (rial * 1e6)).round(2).values
        out[f"معدل {p} (میلیون تومان)"] = (bal[p] / (rial * 1e6)).round(2).values
    if len(periods) >= 2:
        P, L = periods[-2], periods[-1]
        gg = amt[L] / amt[P].replace(0, np.nan) - 1
        churn = (amt[L] == 0) & (amt[P] > 0)
        gg[churn] = -1
        out["وضعیت روند"] = trend_label(churn, gg, cfg["stable_band"], cfg["drop_pct"])
    return out.sort_values(f"مبلغ {periods[-1]} (میلیون تومان)", ascending=False).reset_index(drop=True)


# ------------------------------------------------------------------
# نوشتن اکسل
# ------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def style_sheet(ws, df, start_row=1):
    ws.sheet_view.rightToLeft = True
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, val in enumerate(row, start=1):
            if isinstance(val, float) and (np.isnan(val) or np.isinf(val)):
                val = None
            ws.cell(row=r_idx, column=c_idx, value=val)
    hdr = start_row
    for c_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=hdr, column=c_idx)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(c_idx)
        ws.column_dimensions[letter].width = max(14, min(40, len(str(col)) + 4))
        fmt = None
        if "%" in col:
            fmt = "0.0%"
        elif "میلیون" in col:
            fmt = "#,##0.00"
        elif "تراکنش" in col or "تعداد" in col or "نمره" in col:
            fmt = "#,##0"
        if fmt:
            for r in range(hdr + 1, hdr + 1 + len(df)):
                ws[f"{letter}{r}"].number_format = fmt
    ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    ws.row_dimensions[hdr].height = 42


def write_sheet(wb, name, df):
    ws = wb.create_sheet(name[:31])
    if df.empty:
        ws["A1"] = "موردی یافت نشد"
        ws.sheet_view.rightToLeft = True
        return ws
    style_sheet(ws, df)
    return ws


def add_group_charts(wb, group_name, pack, periods):
    ws = wb.create_sheet(group_name[:31])
    ws.sheet_view.rightToLeft = True
    ws["A1"] = f"{group_name} - روند معدل حساب و تعداد تراکنش (هر سری = یک مشتری)"
    ws["A1"].font = Font(bold=True, size=13)
    n = len(pack["custs"])
    if n == 0:
        ws["A3"] = "موردی یافت نشد"
        return

    def block(top_row, title, values, unit):
        ws.cell(row=top_row, column=1, value=title).font = Font(bold=True)
        hdr = top_row + 1
        ws.cell(row=hdr, column=1, value="مشتری")
        for j, p in enumerate(periods, start=2):
            ws.cell(row=hdr, column=j, value=p)
        for i, (name, vals) in enumerate(zip(pack["custs"], values), start=hdr + 1):
            ws.cell(row=i, column=1, value=name)
            for j, v in enumerate(vals, start=2):
                ws.cell(row=i, column=j, value=v)
        chart = LineChart()
        chart.title = f"{title} - {group_name}"
        chart.y_axis.title = unit
        chart.x_axis.title = "دوره"
        chart.height, chart.width = 9, 18
        data = Reference(ws, min_col=1, max_col=1 + len(periods), min_row=hdr + 1, max_row=hdr + n)
        chart.add_data(data, from_rows=True, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, max_col=1 + len(periods), min_row=hdr))
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.legend.position = "b"
        return chart, hdr + n

    c1, end1 = block(3, "معدل حساب", pack["bal"], "میلیون تومان")
    c2, _ = block(end1 + 3, "تعداد تراکنش", pack["cnt"], "تعداد تراکنش")
    ws.add_chart(c1, "H3")
    ws.add_chart(c2, "H24")
    ws.column_dimensions["A"].width = 38


def write_report(path, res, ipg_df, cfg):
    wb = Workbook()
    wb.remove(wb.active)

    # خلاصه
    s = res["summary"]
    ws = wb.create_sheet("خلاصه")
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "گزارش پایش و پرستاری پایانه‌های فروش"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "دوره‌های بررسی‌شده"
    ws["B3"] = "، ".join(s["periods"])
    row = 5
    for label, val in s["rows"]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=val)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="آمار هر دوره").font = Font(bold=True)
    row += 1
    for j, h in enumerate(["دوره", "تعداد مشتری", "تعداد پایانه"], start=1):
        ws.cell(row=row, column=j, value=h).font = Font(bold=True)
    row += 1
    pos_rows = s["period_stats"]
    for p in s["periods"]:
        ws.cell(row=row, column=1, value=p)
        ws.cell(row=row, column=2, value=int(pos_rows.at[p, "c"]))
        ws.cell(row=row, column=3, value=int(pos_rows.at[p, "t"]))
        row += 1
    row += 1
    notes = [
        "نکات مهم کیفیت داده:",
        f"• فایل {s['L']} فقط یک پایانه برای هر مشتری دارد؛ به همین دلیل مقایسه تغییرات روی «پایانه‌های مشترک» دو دوره انجام شد.",
        f"• تحلیل چندپوز با دوره {s['D']} انجام شده، چون آخرین دوره‌ای است که جزئیات چند پایانه‌ای دارد.",
        "• پایانه‌های غایب (در دوره‌های قبل بوده ولی در دوره آخر نیستند) به‌عنوان احتمال خرابی دستگاه در فهرست تماس آمده‌اند؛ اگر فایل آخر دوباره با همه پایانه‌ها صادر شود، بخشی از این موارد حذف می‌شود.",
        f"• مبالغ داده‌ها ریال فرض شده‌اند و در گزارش به میلیون تومان نمایش داده می‌شوند (هر تومان = {cfg['rial_per_toman']} ریال).",
        f"• دوره‌ها در فایل‌ها یکسان نیستند ({', '.join(s['periods'])})؛ «ریزش» یعنی نبود تراکنش در دوره آخر نسبت به دوره‌های قبل.",
    ]
    for t in notes:
        ws.cell(row=row, column=1, value=t)
        row += 1
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 30

    # جدول‌ها
    write_sheet(wb, "فهرست_تماس", res["calls"])
    write_sheet(wb, "روند_مشتریان", res["customers"])
    write_sheet(wb, "ریزش", res["churn"])
    write_sheet(wb, "کاندید_صبا", res["saba"])
    write_sheet(wb, "پتانسیل", res["potential"])
    write_sheet(wb, "چند_پوز", res["multi"])
    write_sheet(wb, "وضعیت_پایانه‌ها", res["terminals"])
    # نمودارها
    for name, pack in res["charts"].items():
        add_group_charts(wb, name, pack, s["periods"])
    if ipg_df is not None and not ipg_df.empty:
        write_sheet(wb, "درگاه_IPG", ipg_df)

    # تنظیمات
    cfg_df = pd.DataFrame([(k, v) for k, v in cfg.items()], columns=["پارامتر", "مقدار"])
    write_sheet(wb, "تنظیمات", cfg_df)

    wb.save(path)


# ------------------------------------------------------------------
def run(files, out):
    """اجرای کامل تحلیل روی لیست فایل‌ها و نوشتن گزارش در مسیر out."""
    pos_list, ipg_list = [], []
    for f in files:
        kind, period, df = read_snapshot(f)
        if kind == "POS":
            pos_list.append(parse_pos(df, period))
        elif kind == "IPG":
            ipg_list.append(parse_ipg(df, period))
        else:
            print(f"[نادیده گرفته شد] {f}")

    if not pos_list:
        raise SystemExit("هیچ فایل POS پیدا نشد.")
    pos = pd.concat(pos_list, ignore_index=True)
    res = analyse_pos(pos, CONFIG)
    ipg_df = analyse_ipg(pd.concat(ipg_list, ignore_index=True), CONFIG) if ipg_list else None
    write_report(out, res, ipg_df, CONFIG)

    print("گزارش ساخته شد:", out)
    for label, val in res["summary"]["rows"]:
        print(f"  {label}: {val}")


def main():
    ap = argparse.ArgumentParser(description="تحلیل پایش پایانه‌های فروش")
    ap.add_argument("files", nargs="+", help="فایل‌های اکسل ماهانه (POS و/یا IPG)")
    ap.add_argument("-o", "--out", default="pos_monitoring_report.xlsx")
    args = ap.parse_args()
    run(args.files, args.out)


if __name__ == "__main__":
    main()
